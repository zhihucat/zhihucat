from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "蓝血互动游戏UI实现说明.docx"
ASSETS = ROOT / "frontend" / "public" / "assets" / "kanshan"

INK = "171717"
BLUE = "315CFF"
DARK_BLUE = "1E3676"
LIGHT_BLUE = "EAF0FF"
PALE_BLUE = "F5F7FF"
PAPER = "F3EEE3"
PAPER_BRIGHT = "FFFDF7"
MUTED = "D9D1C2"
ORANGE = "F05A28"
YELLOW = "F2CB45"
GREEN = "407A57"
PALE_GREEN = "E8F4E9"
RED = "B93636"
PALE_RED = "FCEAEA"
GRAY = "625C53"
WHITE = "FFFFFF"

# LibreOffice's headless macOS renderer can enumerate Hiragino Sans GB but
# currently emits missing-glyph boxes for it. Arial Unicode MS is installed on
# the same system, covers simplified Chinese, and renders consistently in both
# LibreOffice and Microsoft Word.
BODY_FONT = "Arial Unicode MS"
MONO_FONT = "Arial"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color: str = MUTED, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True) -> None:
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines
    p_pr = paragraph._p.get_or_add_pPr()
    widow = p_pr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        p_pr.append(widow)
    widow.set(qn("w:val"), "1")


def set_no_cell_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_picture_alt(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("蓝血互动游戏 UI 实现说明   ")
    set_run_font(run, size=8.5, color="6D685F")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def paragraph(doc, text="", style=None, *, bold_lead=None, color=None, size=None, align=None, space_after=5, line=1.25):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True, color=color or INK, size=size)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, color=color or INK, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color or INK, size=size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line
    set_paragraph_keep(p)
    return p


def bullet(doc, text, level=0, color=None, checkbox=False):
    p = doc.add_paragraph(style="Checklist" if checkbox else "Spec Bullet")
    p.paragraph_format.left_indent = Inches(0.23 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(("□ " if checkbox else "• ") + text)
    set_run_font(r, size=10.1, color=color or INK)
    set_paragraph_keep(p)
    return p


def add_heading(doc, text, level=1, *, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    set_run_font(r, size={1: 18, 2: 13.5, 3: 11.5}[level], bold=True, color=INK)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_status(doc, label: str, text: str, color: str) -> None:
    p = doc.add_paragraph(style="Status Line")
    r = p.add_run(label)
    set_run_font(r, size=9.2, bold=True, color=color)
    r = p.add_run("  " + text)
    set_run_font(r, size=9.8, color=INK)
    set_paragraph_keep(p)


def add_table(doc, headers, rows, widths=None, header_fill=DARK_BLUE, font_size=8.9, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_no_cell_split(hdr)
    for i, heading in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell, 105, 115, 105, 115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(str(heading))
        set_run_font(r, size=font_size, bold=True, color=WHITE if header_fill in (DARK_BLUE, INK) else INK)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_no_cell_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_border(cell)
            set_cell_margins(cell, 95, 110, 95, 110)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(str(value)) <= 10 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=bool(first_col_bold and i == 0), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_small_label(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, name=MONO_FONT, size=8.2, bold=True, color=BLUE)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_figure(doc, path: Path, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(shape, caption, alt)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8.5, color=GRAY)
    set_paragraph_keep(cap)


_PAGE_BREAK_COUNT = 0


def page_break(doc):
    """Keep the cover separate; let later sections flow to avoid sparse pages."""
    global _PAGE_BREAK_COUNT
    if _PAGE_BREAK_COUNT == 0:
        doc.add_page_break()
    _PAGE_BREAK_COUNT += 1


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.67)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_after = Pt(14)
    title_p_pr = title._element.get_or_add_pPr()
    old_title_border = title_p_pr.find(qn("w:pBdr"))
    if old_title_border is not None:
        title_p_pr.remove(old_title_border)
    title_border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        title_border.append(element)
    title_p_pr.append(title_border)

    for level, size in ((1, 18), (2, 13.5), (3, 11.5)):
        style = styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6 if level == 1 else 4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = BODY_FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.font.size = Pt(8.5)

    if "Spec Bullet" not in styles:
        bullet_style = styles.add_style("Spec Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet_style = styles["Spec Bullet"]
    bullet_style.font.name = BODY_FONT
    bullet_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    bullet_style.font.size = Pt(10.1)
    bullet_style.paragraph_format.space_after = Pt(3.5)
    bullet_style.paragraph_format.line_spacing = 1.22

    if "Checklist" not in styles:
        checklist_style = styles.add_style("Checklist", WD_STYLE_TYPE.PARAGRAPH)
    else:
        checklist_style = styles["Checklist"]
    checklist_style.font.name = BODY_FONT
    checklist_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    checklist_style.font.size = Pt(10)
    checklist_style.paragraph_format.space_after = Pt(4)
    checklist_style.paragraph_format.line_spacing = 1.2

    if "Status Line" not in styles:
        status_style = styles.add_style("Status Line", WD_STYLE_TYPE.PARAGRAPH)
    else:
        status_style = styles["Status Line"]
    status_style.font.name = BODY_FONT
    status_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    status_style.font.size = Pt(9.8)
    status_style.paragraph_format.space_after = Pt(4)
    status_style.paragraph_format.line_spacing = 1.2

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(44)
    r = p.add_run("ZHIHU HACKATHON   PRODUCT UI HANDOFF")
    set_run_font(r, name=MONO_FONT, size=9.3, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    r = p.add_run("蓝血互动游戏 UI 实现说明")
    set_run_font(r, size=31, bold=True, color=INK)
    p.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("面向 UI 设计和前端实现的完整页面规范")
    set_run_font(r, size=14.5, bold=True, color=GRAY)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_cell_width(table.cell(0, 0), 4.35)
    set_cell_width(table.cell(0, 1), 2.0)
    for c in table.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c, 0, 0, 0, 0)
        tc_pr = c._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil")
            borders.append(e)
        tc_pr.append(borders)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(8)
    lines = [
        ("产品", "蓝血疑云 单故事互动案卷"),
        ("交付对象", "UI 设计同学和前端实现同学"),
        ("版本", "1.0"),
        ("日期", "2026 年 9 月 14 日"),
        ("核心结论", "当前代码已形成可演示的完整主流程。专职 UI 同学应优先统一视觉层级、核验移动端和状态可读性，再补完成回执。"),
    ]
    for idx, (label, value) in enumerate(lines):
        if idx:
            p = table.cell(0, 0).add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.28
        r = p.add_run(label + "  ")
        set_run_font(r, size=9.5, bold=True, color=BLUE)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)
    portrait = table.cell(0, 1).paragraphs[0]
    portrait.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    shape = portrait.add_run().add_picture(str(ASSETS / "portrait.jpg"), width=Inches(1.92))
    set_picture_alt(shape, "刘看山立姿", "白色背景上的刘看山三维角色立姿参考图")

    paragraph(doc, "本说明把当前代码中已经存在的界面与建议新增项分开标记。实现时不得把知乎公开片段之外的推演写成原作事实，也不得把接口失败伪装成已核验原文。", size=10.5, color=GRAY, space_after=0, line=1.45)

    page_break(doc)

    add_heading(doc, "一 文档用途与交付结论", 1)
    paragraph(doc, "本说明用于把《蓝血》互动案卷交给 UI 设计和前端实现同学。范围覆盖入口、刘看山故事向导、故事开场、三场景搜证、来源阅读器、证据组牌、解释力对决、失败、阶段推演 End、完成后的回执建议，以及所有跨页面状态。")
    paragraph(doc, "当前产品的可交付核心是一条单故事闭环：玩家理解故事边界，找到六条线索，组成四张证据牌，挑战日常解释，再选择阶段假说和知乎求证问题。UI 的首要任务是让玩家随时知道自己看到的是事实、策划转述还是推断。", bold_lead="当前产品的可交付核心")

    add_heading(doc, "状态标记", 2)
    add_status(doc, "已实现", "当前代码已有可见界面与交互，可进入视觉精修和回归验收。", GREEN)
    add_status(doc, "已接入待联调", "界面路径已存在，但在线结果依赖知乎接口与运行环境，必须同时保留透明降级。", ORANGE)
    add_status(doc, "建议新增", "文档提出的下一步方案，当前代码未实现，不应出现在完成清单中。", BLUE)
    add_status(doc, "内容边界", "属于长期约束，不得以动效、文案或结果页暗示原作真相。", RED)

    add_heading(doc, "当前界面清单", 2)
    add_table(doc,
        ["界面", "当前状态", "本轮 UI 重点"],
        [
            ["全局顶栏", "已实现", "品牌、语言、玩家、积分与账户操作的密度和移动端折叠"],
            ["故事入口页", "已实现", "故事钩子、来源状态、试玩时长和主按钮的首屏优先级"],
            ["刘看山首页导览", "已实现", "三主题切换、角色占比、长文阅读节奏"],
            ["刘看山悬浮向导", "已实现", "上下文提示、遮挡控制和移动端抽屉体验"],
            ["故事开场", "已实现", "角色任务、信息来源和开始搜证动作"],
            ["三场景搜证", "已实现", "六条线索发现状态和三栏布局"],
            ["来源阅读器", "已接入待联调", "原文短引、策划转述、离线降级必须显式区分"],
            ["六选四组牌", "已实现", "三张固定锚点与一张方向牌的规则提示"],
            ["解释力对决和失败", "已实现", "战斗数值、出牌反馈、重试恢复路径"],
            ["阶段推演 End", "已实现", "假说、证据缺口、下一步验证和求证问题"],
            ["完成回执", "建议新增", "完成后保留本局结论和可复制内容，再返回入口页"],
            ["移动端实机验收", "待完成", "代码已有响应式规则，仍需覆盖小屏、软键盘和安全区"],
        ], widths=[1.48, 1.1, 4.45], font_size=8.8, first_col_bold=True)

    add_heading(doc, "文档使用方式", 2)
    bullet(doc, "UI 设计同学先按页面章节完成高保真与组件状态，不另行改写产品规则。")
    bullet(doc, "前端同学按数据状态映射实现，接口来源标识优先于视觉装饰。")
    bullet(doc, "联调时先验收 P0 清单，再决定是否投入建议新增的 P1 和 P2。")

    page_break(doc)

    add_heading(doc, "二 完整体验流程", 1)
    paragraph(doc, "主流程只有一条，失败会回到搜证与组牌，不新开支线。每一步都必须显示当前阶段、下一步动作和离开条件。")
    add_table(doc,
        ["阶段", "玩家目标", "主要动作", "离开条件", "刘看山职责"],
        [
            ["入口页", "理解故事和内容边界", "查看来源状态并开始", "点击进入调查", "给出一分钟背景导览"],
            ["故事开场", "确认身份与任务", "阅读简介和来源说明", "点击开始搜证", "不覆盖模态框"],
            ["线索搜证", "找齐三场景六条线索", "切换场景并点击热点", "发现进度达到六条", "解释事实与推断的差别"],
            ["证据组牌", "形成四张可质证牌组", "选择三张锚点和一张方向牌", "四张齐全且锚点完整", "提示第四张牌的取向"],
            ["解释力对决", "让日常解释无法覆盖全部证据", "依次出示四张证据并管理体力", "解释力归零且四张均出示", "解释反方与证据链"],
            ["失败", "理解本轮链条缺口", "查看分析后返回重组", "回到搜证状态", "强调失败不等于假说错误"],
            ["阶段推演", "选择当前假说和求证问题", "比较四个假说和三个问题", "两项均已选择", "说明推荐依据与证据缺口"],
            ["完成", "保存结果并决定下一步", "当前为封存并返回入口", "成绩写入本地或账户", "建议增加完成回执"],
        ], widths=[0.75, 1.35, 1.55, 1.35, 2.05], font_size=8.45, first_col_bold=True)

    add_heading(doc, "阶段门槛", 2)
    bullet(doc, "进入质证前必须发现全部六条线索。")
    bullet(doc, "牌组固定为四张，包括指尖红血、同事蓝转红、定向试卷三张关键锚点，以及城市资料或灰夹克之一。")
    bullet(doc, "培训教材是背景线索，可以查阅但不入庭。")
    bullet(doc, "对决中每张入选证据至少实际出示一次；重复强牌不能绕过证据链要求。")
    bullet(doc, "End 必须同时选择一个阶段假说和一个知乎求证问题，完成按钮才可用。")

    add_heading(doc, "内容判断规则", 2)
    add_table(doc,
        ["信息层", "可用表达", "禁止表达"],
        [
            ["事实", "公开片段直接写出的观察或本次接口即时命中的短引", "把策划复述写成逐字原文"],
            ["推断", "世界置换、记忆错位、观察实验等竞争假说", "正确答案是 或 幕后黑手就是"],
            ["证据缺口", "目前仍不能排除什么以及下一步如何验证", "以作者暗示为由补写机制"],
            ["阶段结果", "本局证据链更支持某种解释", "解锁原作结局 或 宣称还原真相"],
        ], widths=[0.9, 3.1, 3.05], font_size=8.8, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "三 视觉语言与设计 Token", 1)
    paragraph(doc, "现有界面采用纸张底色、黑色粗边框和高饱和蓝黄强调，接近可触摸的案卷卡片。UI 精修应保持这一辨识度，同时降低同屏元素竞争，优先保证信息层级。")

    add_heading(doc, "颜色", 2)
    add_table(doc,
        ["Token", "色值", "用途"],
        [
            ["Ink", "#171717", "正文、标题、主边框和主要阴影"],
            ["Paper", "#F3EEE3", "页面底色"],
            ["Paper Bright", "#FFFDF7", "面板、卡片和对话区"],
            ["Paper Muted", "#D9D1C2", "弱边框、禁用分隔和未激活步骤"],
            ["Primary Blue", "#315CFF", "主按钮、激活状态、焦点和来源强调"],
            ["Periwinkle", "#8EA2ED", "弱选中、辅助背景和虚线提示"],
            ["Orange", "#F05A28", "风险提示和过程反馈"],
            ["Yellow", "#F2CB45", "教程提示、计时与轻量强调"],
            ["Green", "#407A57", "完成、在线或可通过状态"],
            ["Red", "#B93636", "错误、失败和高风险状态"],
        ], widths=[1.3, 1.15, 4.6], font_size=8.8, first_col_bold=True)

    add_heading(doc, "字体与层级", 2)
    add_table(doc,
        ["层级", "建议规格", "用途"],
        [
            ["展示标题", "32 至 76 px 900 字重", "入口主标题，响应式缩放"],
            ["页面标题", "24 至 38 px 850 字重", "故事开场与 End 主标题"],
            ["组件标题", "16 至 22 px 800 字重", "面板、场景和结果标题"],
            ["正文", "14 至 16 px 400 或 600 字重", "故事说明和交互解释"],
            ["状态与数据", "10 至 12 px 等宽字体", "标签、计数、来源状态和 Work ID"],
        ], widths=[1.25, 2.2, 3.6], font_size=8.8, first_col_bold=True)
    paragraph(doc, "当前 CSS 使用 Inter、系统无衬线字体和 SFMono 等宽字体。中文版要显式提供中文系统字体回退，避免数字与中文基线漂移。", size=9.8, color=GRAY)

    add_heading(doc, "结构与动效", 2)
    add_table(doc,
        ["项目", "当前基线", "实现要求"],
        [
            ["间距", "以 4 和 8 px 为基础", "优先使用 4 8 12 16 24 32，避免零散值继续扩散"],
            ["边框", "关键卡片 3 px，次级控件 2 px", "状态变化时保留轮廓，不只改变底色"],
            ["圆角", "4 至 12 px", "按钮小圆角，模态框与角色卡使用 10 至 12 px"],
            ["阴影", "3 至 7 px 硬阴影", "只用于可交互或层级提升，不给所有文本块加阴影"],
            ["过渡", "约 150 至 250 ms", "位移不超过 3 px，战斗反馈除外"],
            ["焦点", "3 px 蓝色外框", "键盘焦点不可被 hover 或阴影覆盖"],
        ], widths=[1.0, 2.05, 4.0], font_size=8.8, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "四 全局导航与通用组件", 1)

    add_heading(doc, "全局顶栏", 2)
    add_status(doc, "已实现", "桌面端为吸顶栏，左侧品牌，右侧依次为中英文切换、玩家入口、积分和账户动作。", GREEN)
    bullet(doc, "品牌按钮返回故事入口页。图标、名称和副标题视为一个点击目标。")
    bullet(doc, "语言切换使用 tab 语义并保留 aria selected。切换后页面标题和 html lang 同步。")
    bullet(doc, "玩家入口显示头像、昵称或登录提示；资料加载时禁用，失败时出现非阻塞同步说明。")
    bullet(doc, "移动端可隐藏次要文字，但账号和语言入口不能被刘看山悬浮按钮遮挡。")

    add_heading(doc, "通用组件", 2)
    add_table(doc,
        ["组件", "默认形态", "关键状态"],
        [
            ["主按钮", "蓝底黑边与硬阴影", "hover 轻微上移，disabled 降低不透明度并保留文本"],
            ["次按钮", "亮纸底黑边", "返回、查看来源、重新组牌"],
            ["面板", "亮纸底 3 px 黑边 10 px 圆角", "通过标题、徽标和内容区建立层级"],
            ["标签", "等宽小字", "ready、来源模式、关键证据和风险等级"],
            ["阶段轨道", "搜证 解释力对决 END", "active 蓝底，done 绿底，未开始为浅边框"],
            ["来源条", "蓝色左边框", "标题、作者、Work ID、来源通知"],
            ["错误", "红色边框或红色文字", "role alert，紧邻失败动作，并提供重试"],
            ["空状态", "弱背景与简短说明", "只说明下一步，不展示技术错误详情"],
        ], widths=[1.1, 2.35, 3.6], font_size=8.8, first_col_bold=True)

    add_heading(doc, "来源状态统一写法", 2)
    bullet(doc, "在线且存在命中短引时显示 知乎 API 已连接，并在每条命中材料中显示 已核验原文短引。")
    bullet(doc, "在线但锚点未命中时显示 原文锚点未命中，内容只标 策划转述。")
    bullet(doc, "接口失败时显示 策划离线模式 和 透明降级，不使用断网错误页阻断试玩。")
    bullet(doc, "所有模式都保留故事名、作者、Work ID 和版权边界。")

    add_heading(doc, "组件实现来源", 2)
    add_table(doc,
        ["文件", "职责"],
        [
            ["frontend app page tsx", "主流程、场景、组牌、对决、失败和 End"],
            ["frontend app components kanshan guide tsx", "首页导览和上下文悬浮向导"],
            ["frontend app components cat evidence tsx", "来源阅读器与线索展示"],
            ["frontend app globals css", "全部视觉 token、布局、动效和响应式规则"],
            ["frontend app lib court battle ts", "手牌、体力、护盾、计时和胜负状态"],
            ["backend src zhihu story js", "六条线索、来源模式、假说和求证问题"],
        ], widths=[2.65, 4.4], font_size=8.6, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "五 故事入口页", 1)
    add_status(doc, "已实现", "入口页已经形成故事钩子、来源状态、核心数字、内容边界、开始按钮、刘看山导览和四步流程。", GREEN)

    add_heading(doc, "页面结构", 2)
    add_table(doc,
        ["区域", "必须显示", "交互与状态"],
        [
            ["顶部标题", "蓝血疑云 谁的世界出了错", "标题下只用一句话解释玩法，不堆设定"],
            ["来源状态", "在线或离线 Work ID 或降级原因", "在线用绿点，离线也要明确可继续试玩"],
            ["故事主卡", "封面、标签、故事钩子", "远程封面不可用时使用 BLUE BLOOD 文本占位"],
            ["数据概览", "3 场景 6 线索 4 卡牌 1 求证问题", "作为玩法承诺，数字不可与数据不一致"],
            ["内容边界", "仅用公开片段 阶段推演不代表原作结局", "始终在主按钮附近可见"],
            ["主动作", "封存剧情进入调查 或 重新进入调查", "完成后显示最佳裁决分数"],
            ["故事向导", "刘看山大卡片", "三主题可切换，另有进入搜证按钮"],
            ["流程条", "阅读简介 原文搜证 六选四组牌 选择求证", "说明规则，不承担导航"],
        ], widths=[1.15, 2.55, 3.35], font_size=8.65, first_col_bold=True)

    add_heading(doc, "桌面布局", 2)
    bullet(doc, "入口内容宽度上限约 1180 px，标题区为主信息加来源状态双列。")
    bullet(doc, "故事主卡左侧封面，右侧钩子、数字和主动作。主按钮必须进入首屏可见范围。")
    bullet(doc, "刘看山导览位于故事主卡之后，角色区与正文区视觉分栏。")

    add_heading(doc, "移动端布局", 2)
    bullet(doc, "标题与来源状态改为单列；封面、内容和动作纵向排列。")
    bullet(doc, "四项数据可使用两列网格，不压缩到一行。")
    bullet(doc, "开始按钮和刘看山导览按钮均使用全宽，按钮间至少 8 px。")
    bullet(doc, "标签超过一行时自然换行，不横向滚动。")

    add_heading(doc, "入口验收", 2)
    bullet(doc, "断网状态下玩家仍能进入完整故事，且不会误以为内容来自实时接口。", checkbox=True)
    bullet(doc, "完成过一次后，入口按钮和最佳分数状态立即更新。", checkbox=True)
    bullet(doc, "刘看山大卡片不会把主故事 CTA 挤出第一屏以下过远。", checkbox=True)

    page_break(doc)

    add_heading(doc, "六 刘看山故事向导", 1)
    paragraph(doc, "刘看山是知乎故事向导，不是《蓝血》角色，也不是自由对话机器人。当前实现使用确定性预设文案，目的在于减少幻觉和演示不确定性。")

    add_figure(doc, ASSETS / "still.png", 1.35, "刘看山静态降级素材", "刘看山站立的透明背景静态图，用于减少动态效果模式")

    add_heading(doc, "首页大卡片", 2)
    add_status(doc, "已实现", "角色在左，内容在右；移动端上下堆叠。默认打开故事发生了什么。", GREEN)
    add_table(doc,
        ["主题", "内容目标", "结束动作"],
        [
            ["故事发生了什么", "交代蓝血常识、方诺红血、同事蓝转红、公共记录、定向测试与跟随者", "继续切换主题或开始搜证"],
            ["我在游戏里做什么", "说明三场景六线索、三张锚点加一张方向牌、阶段推演与知乎问题", "开始搜证"],
            ["哪些还不能确定", "说明世界置换、记忆错位与观察实验都只是阶段假说", "开始搜证"],
        ], widths=[1.45, 4.45, 1.15], font_size=8.7, first_col_bold=True)
    bullet(doc, "切换主题时只替换答案区和角色动作，不改变整卡高度，避免页面跳动。")
    bullet(doc, "标签使用 tab 语义；答案区使用 aria live polite，切换后不抢焦点。")
    bullet(doc, "首页角色动作仅显示一个 GIF。减少动态效果模式显示 still png。")

    add_heading(doc, "悬浮向导", 2)
    add_status(doc, "已实现", "调查、对决和 End 均提供右下角 问看山 入口；关闭后只保留胶囊触发按钮。", GREEN)
    add_table(doc,
        ["上下文", "默认提示", "可追问内容"],
        [
            ["急救培训室", "比较教材 方诺 同事三个独立观察", "事实还是推断 证据强度 第四张牌"],
            ["深夜资料库", "寻找公共记录与只针对方诺的安排", "事实还是推断 证据强度 第四张牌"],
            ["老街死胡同", "区分被跟踪的感受与观察到的行为", "事实还是推断 证据强度 第四张牌"],
            ["离线搜证", "当前材料是策划转述", "为什么不能称为原文"],
            ["对决进行中", "反方代表疲劳 误记和巧合", "质疑对象 证据链 重复出牌"],
            ["对决胜利", "四张证据均已出示", "为何能进入阶段推演"],
            ["对决失败", "本轮链条未闭合", "如何重组和保护论证"],
            ["阶段推演", "选择下一步验证而非原作答案", "推荐依据 其他假说 安全提问"],
        ], widths=[1.28, 3.15, 2.62], font_size=8.5, first_col_bold=True)

    add_heading(doc, "交互规范", 2)
    bullet(doc, "桌面端面板宽度上限约 400 px，触发器固定在安全区内；展开后先显示情境提示。")
    bullet(doc, "移动端面板宽度为视口减 1.2 rem，高度上限约 64 d vh，内部滚动。")
    bullet(doc, "任何时候都不提供自由文本输入。若未来接入大模型，必须另设来源约束、拒答和兜底，不属于当前实现。")
    bullet(doc, "禁止使用 正确答案是 真相是 作者暗示了 幕后黑手就是 解锁原作结局。")

    page_break(doc)

    add_heading(doc, "七 故事开场与调查框架", 1)

    add_heading(doc, "故事开场", 2)
    add_status(doc, "已实现", "进入故事后先出现模态开场卡，包含来源模式、封面、简介、玩家身份、任务和版权说明。", GREEN)
    add_table(doc,
        ["元素", "内容", "实现要求"],
        [
            ["来源模式", "实时接口 或 离线降级", "位于卡片顶部并保持可读"],
            ["故事简介", "方诺经历与公开片段范围", "不提前展示假说答案"],
            ["玩家身份", "方诺的匿名调查协助者", "用第二人称解释目标"],
            ["任务", "找六条线索 三张锚点加一张方向牌", "黄色任务区突出完成门槛"],
            ["来源署名", "知乎故事名 作者和通知", "保持在按钮之前"],
            ["动作", "返回故事页 开始搜证", "开始为主按钮，返回为次按钮"],
        ], widths=[1.1, 2.65, 3.3], font_size=8.8, first_col_bold=True)
    add_status(doc, "建议新增", "为模态框补齐焦点锁定、Esc 关闭策略和关闭后的焦点恢复。当前已有 dialog 与 aria modal 语义，但未见完整焦点管理。", BLUE)

    add_heading(doc, "调查页全局框架", 2)
    bullet(doc, "顶部紧凑导航包含返回按钮和三阶段轨道。")
    bullet(doc, "案件摘要区显示当前阶段、故事标题、目标、已取证数、关键证据数和总分。")
    bullet(doc, "来源条显示故事名、作者、Work ID 和在线或离线说明。")
    bullet(doc, "桌面端主体为故事现场、来源阅读器、证据牌组三栏；每栏独立滚动，页面本身尽量不产生双重滚动。")
    bullet(doc, "刘看山悬浮向导持续存在，但不得遮住进入质证按钮和 End 操作区。")

    add_heading(doc, "调查交互", 2)
    add_table(doc,
        ["动作", "立即反馈", "连带变化"],
        [
            ["切换场景", "选中标签高亮", "阅读器与已发现证据保持不变"],
            ["点击未发现热点", "热点变为已收集", "右侧出现证据卡 进度加一 得分加五"],
            ["再次点击已发现热点", "热点恢复未收集", "证据卡和已选状态移除 得分减五"],
            ["打开来源材料", "文档按钮激活", "阅读器显示短引或策划转述"],
            ["选择证据卡", "卡片出现已选状态", "上庭计数增加，不移出证据库"],
            ["取消证据卡", "卡片恢复默认", "上庭计数减少，调查发现状态保留"],
        ], widths=[1.55, 2.45, 3.05], font_size=8.7, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "八 三场景六线索搜证", 1)
    paragraph(doc, "六条线索按叙事功能分为共同背景、关键锚点和推理方向。视觉必须让玩家先搜齐，再理解哪些能上庭。")

    add_table(doc,
        ["场景", "线索", "类型", "可信度", "牌组角色"],
        [
            ["急救培训室", "培训教材的蓝血常识", "书证", "9", "共同背景 不入庭"],
            ["急救培训室", "指尖流出的红血", "影像物证", "8", "关键锚点 必选"],
            ["急救培训室", "同事由蓝变红的血迹", "影像物证", "9", "关键锚点 必选"],
            ["深夜资料库", "与记忆冲突的城市资料", "书证", "8", "方向牌 世界规则置换"],
            ["深夜资料库", "只针对方诺的测试卷", "书证", "9", "关键锚点 必选"],
            ["老街死胡同", "没有折返的灰夹克", "影像物证", "7", "方向牌 主动观察"],
        ], widths=[1.25, 2.35, 1.0, 0.72, 1.73], font_size=8.75, first_col_bold=True)

    add_heading(doc, "场景一 急救培训室", 2)
    paragraph(doc, "目标是建立三种独立观察：教材与同事共享蓝血常识，方诺本人从一开始流出红血，同事的血实际由蓝转红。热点卡可用血液和培训素材，但不能用惊悚血腥画面压过推理。")
    add_heading(doc, "场景二 深夜资料库", 2)
    paragraph(doc, "目标是把异常从身体扩展到公共记录与定向筛查。城市资料应表现为系统化档案，不用模糊故障特效；测试卷则要清楚呈现只有方诺拿到常识题。")
    add_heading(doc, "场景三 老街死胡同", 2)
    paragraph(doc, "目标是把被跟踪的感觉转化为可观察行为。画面可使用长巷和消失路径，但文案必须保留误认、藏身等日常解释。")

    add_heading(doc, "热点与证据卡状态", 2)
    add_table(doc,
        ["状态", "热点表现", "证据卡表现"],
        [
            ["未发现", "显示调查提示与线索插画", "不出现在证据库"],
            ["已发现", "显示已收集和勾选", "出现标题 证明目的 描述和可信度"],
            ["可入庭未选", "保持已发现", "可点击并显示 点击选择带上庭"],
            ["已选入庭", "保持已发现", "选中边框与 已选入庭 点击取消"],
            ["背景锁定", "保持已发现", "显示背景线索，不允许入庭"],
        ], widths=[1.25, 2.55, 3.5], font_size=8.75, first_col_bold=True)

    add_heading(doc, "视觉验收重点", 2)
    bullet(doc, "六条热点在 1366 px 宽桌面上无需横向滚动即可完成。", checkbox=True)
    bullet(doc, "已发现、已选择和背景锁定至少有文字与轮廓两种提示。", checkbox=True)
    bullet(doc, "场景切换不会清除已发现或已选择证据。", checkbox=True)

    page_break(doc)

    add_heading(doc, "九 来源阅读器与内容标识", 1)
    paragraph(doc, "来源阅读器的核心价值不是模拟阅读全文，而是把每条线索的来源状态讲清楚。当前后端只在实时接口响应中命中预设短引时展示短引；完整正文不落盘。")

    add_heading(doc, "三种来源模式", 2)
    add_table(doc,
        ["数据条件", "页面标题与徽标", "单条材料标记", "允许展示"],
        [
            ["zhihu live 且命中短引", "公开片段锚点 与 N 条短引已核验", "命中项为 查看已核验短引，未命中项仍为 查看策划转述", "即时核验的短引 位置 作者 策划释义"],
            ["zhihu live 但零命中", "策划线索转述 与 原文锚点未命中", "全部为 查看策划转述", "策划转述与未命中说明"],
            ["curated fallback", "离线案卷线索 与 策划离线转述", "全部为 查看策划转述", "策划转述与接口恢复说明"],
        ], widths=[1.42, 2.05, 1.85, 2.0], font_size=8.35, first_col_bold=True)

    add_heading(doc, "阅读器结构", 2)
    bullet(doc, "左上文档列表显示六条案卷材料，每项包含线索插画、标题和来源动作。")
    bullet(doc, "阅读区展示材料标题、类型、正文、证据指针和发现动作。已发现指针必须可逆。")
    bullet(doc, "原文短引使用固定抬头 知乎公开片段 原文短引；策划内容使用 策划线索转述 非原文全文。")
    bullet(doc, "故事名、作者、位置和内容来源必须与短引同屏，不能只放在页面底部。")
    bullet(doc, "外链 查看原作 位于 End，不应把当前阅读器包装成全文阅读器。")

    add_heading(doc, "接口与降级规范", 2)
    add_status(doc, "已接入待联调", "列表与故事详情通过知乎黑客松接口读取；列表或详情失败时返回内置《蓝血》策划案卷。", ORANGE)
    bullet(doc, "接口失败不阻断核心流程，但来源状态必须立刻切换并保留失败原因摘要。")
    bullet(doc, "远程封面只接受知乎与知图 HTTPS 域名；失败时显示稳定占位。")
    bullet(doc, "不得在浏览器本地存储或页面源代码中加入完整正文。")
    bullet(doc, "原文短引命中数量来自 verifiedQuoteCount，不可由 UI 猜测。")

    add_heading(doc, "文案层级", 2)
    add_table(doc,
        ["层级", "示例", "视觉权重"],
        [
            ["来源状态", "知乎 API 已连接 或 策划离线模式", "高 颜色加文字"],
            ["材料级状态", "已核验短引 或 策划转述", "中 紧邻标题"],
            ["内容边界", "完整正文不落盘 推演不代表结局", "中 长期可见"],
            ["技术详情", "Work ID 和接口错误摘要", "低 等宽小字"],
        ], widths=[1.25, 3.3, 2.77], font_size=8.75, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "十 六选四证据组牌", 1)
    add_status(doc, "已实现", "玩家必须先发现六条线索，再组成四张上庭牌。当前可入庭集合为三张关键锚点和两张方向牌，培训教材作为背景不入庭。", GREEN)

    add_heading(doc, "组牌规则", 2)
    add_table(doc,
        ["规则", "UI 呈现", "不满足时"],
        [
            ["六条线索全部发现", "调查 KPI 显示 6，三个场景热点均为已收集", "提示还有几条线索未发现"],
            ["牌组正好四张", "徽标显示 4 比 4，按钮显示带四张牌进入质证", "提示请组成完整四张证据牌组"],
            ["三张关键锚点必选", "关键证据标签与选中边框同时可见", "列出缺少的关键证据标题"],
            ["一张方向牌", "城市资料或灰夹克二选一", "刘看山解释两种推理方向"],
            ["背景线索不可选", "按钮禁用并显示背景线索", "不应只用灰色表达原因"],
            ["最多四张", "计数实时更新", "提示先取消一张"],
        ], widths=[1.45, 3.35, 2.5], font_size=8.65, first_col_bold=True)

    add_heading(doc, "证据卡信息层级", 2)
    bullet(doc, "第一层是标题，必须在两行内读完。")
    bullet(doc, "第二层是证明目的，用箭头提示这张证据能证明什么。")
    bullet(doc, "第三层是事实描述，最多两到三行，详情可在来源阅读器查看。")
    bullet(doc, "底部元信息依次显示证据类型、可信度、关键或支持角色和选择状态。")

    add_heading(doc, "选择反馈", 2)
    add_table(doc,
        ["状态", "边框与背景", "文字与可用性"],
        [
            ["默认", "亮纸底 2 至 3 px 边框", "点击选择带上庭"],
            ["选中", "蓝色或黄色强调并保留黑色轮廓", "已选入庭 点击取消选择"],
            ["关键", "暖黄色系", "显示关键证据"],
            ["方向", "冷蓝色系", "显示支持证据及推理方向"],
            ["背景锁定", "低饱和但保持 4.5 比 1 文字对比", "显示不作为方向牌"],
        ], widths=[1.15, 2.8, 3.35], font_size=8.75, first_col_bold=True)

    add_heading(doc, "进入质证按钮", 2)
    bullet(doc, "按钮文案随选择数更新为 带 N 比 4 张牌进入质证。")
    bullet(doc, "只要至少选一张，按钮可触发并在点击后给出具体错误；视觉设计也可在四张未齐时保留弱提示，但不要让玩家无反馈。")
    bullet(doc, "完成四张规则后进入对决，调查发现和牌组选择保留，失败返回时不要求重新搜证。")

    page_break(doc)

    add_heading(doc, "十一 解释力对决", 1)
    paragraph(doc, "对决不是战胜故事角色，而是用四项材料削弱疲劳、误记和巧合组成的日常解释。所有战斗文案都要保持这一语义。")

    add_heading(doc, "顶部信息条", 2)
    add_table(doc,
        ["区域", "当前数据", "视觉要求"],
        [
            ["左侧", "日常解释力 当前值与最大值", "红色或橙色进度，但必须同时显示数字"],
            ["中间", "回合数 25 秒倒计时 行动阶段", "倒计时低于五秒时增加文字或节奏提示"],
            ["右侧", "论证稳定度 20 专注 10 反驳缓冲 8", "三种资源不要只用相近色条区分"],
        ], widths=[1.0, 3.05, 3.25], font_size=8.7, first_col_bold=True)

    add_heading(doc, "对决舞台", 2)
    bullet(doc, "左侧为待反驳假说，右侧为玩家主张，中间显示最近两轮对话和判定。")
    bullet(doc, "开场案件名只短暂出现约 1.8 秒，不作为永久遮罩。")
    bullet(doc, "出牌后手牌锁定，玩家动画结束后约 2 秒出现反方回应，再约 0.85 秒进入下一轮。")
    bullet(doc, "对话区必须能滚动，不让长文本挤压手牌。")

    add_heading(doc, "手牌与资源", 2)
    add_table(doc,
        ["牌型", "作用", "视觉区别", "限制"],
        [
            ["关键证据", "削弱日常解释力", "暖黄色底与关键标记", "消耗专注 四张都必须至少出示一次"],
            ["方向证据", "削弱日常解释力并决定 End 推荐", "冷蓝色底", "同样属于必须出示的四张"],
            ["体力恢复", "恢复三或五点专注", "绿色或呼吸图形", "满专注时不可用 仍消耗行动"],
            ["防御策略", "增加五或八点反驳缓冲", "蓝灰防御图形", "护盾满时不可用 仍消耗行动"],
        ], widths=[1.15, 2.0, 1.8, 2.35], font_size=8.45, first_col_bold=True)

    add_heading(doc, "胜利条件和反馈", 2)
    bullet(doc, "日常解释力降为零，同时四张入选证据都已经出示，才判定证据链成立。")
    bullet(doc, "重复使用同一张牌最多把日常解释力压到一，不能直接结束。")
    bullet(doc, "胜利后显示 日常解释已被动摇，并提供进入阶段推演按钮。")
    bullet(doc, "音效只跟随已提交的战斗效果；无效点击和等待阶段保持静默。背景音乐进入对决时播放，结束或退出时停止。")
    add_status(doc, "建议新增", "提供显式静音开关并记住选择。当前有背景音乐和反馈音，但没有可见音频控制。", BLUE)

    page_break(doc)

    add_heading(doc, "十二 失败与恢复", 1)
    add_status(doc, "已实现", "论证稳定度降为零时显示证据链未能闭合。玩家可查看本次分析，再返回调查重组。", GREEN)

    add_heading(doc, "失败界面", 2)
    bullet(doc, "标题使用 证据链未能闭合，不使用 真相错误 或 游戏结束。")
    bullet(doc, "说明 日常解释暂时成立，并提示调整出牌顺序或先用防御策略。")
    bullet(doc, "查看本次分析请求阶段裁决，结果页提供返回调查重组。")
    bullet(doc, "返回调查后保留已发现六条线索和已选牌组，重置战斗回合、对话、体力、护盾和结果。")

    add_heading(doc, "错误与失败的区别", 2)
    add_table(doc,
        ["类型", "含义", "用户动作", "视觉"],
        [
            ["玩法失败", "证据链或资源管理未通过", "查看分析并重组", "红色但保持可恢复语气"],
            ["接口错误", "回应或裁决请求失败", "保留当前上下文并重试", "role alert 紧邻操作区"],
            ["来源降级", "知乎接口暂不可用", "继续使用策划案卷", "黄色或蓝色来源提示 非错误阻断"],
            ["规则错误", "线索或牌组不满足门槛", "回到对应卡片补齐", "明确缺少数量或名称"],
        ], widths=[1.2, 2.2, 2.25, 1.65], font_size=8.65, first_col_bold=True)

    add_heading(doc, "恢复体验建议", 2)
    add_status(doc, "建议新增", "在失败分析中高亮本轮未出示的入选证据，并提供 返回并定位牌组 按钮。当前只返回调查顶部。", BLUE)
    add_status(doc, "建议新增", "保留最近一次出牌顺序作为可折叠回顾，帮助玩家理解失败原因。", BLUE)

    add_heading(doc, "失败验收", 2)
    bullet(doc, "失败后不会清空搜证进度。", checkbox=True)
    bullet(doc, "错误提示不会被刘看山面板遮挡。", checkbox=True)
    bullet(doc, "返回调查后可以直接调整第四张方向牌并重新进入。", checkbox=True)
    bullet(doc, "失败文案不暗示原作答案。", checkbox=True)

    page_break(doc)

    add_heading(doc, "十三 阶段推演 End", 1)
    add_status(doc, "已实现", "胜利后进入独立 End 页面。它总结证据链，引导玩家选择阶段假说和知乎求证问题，不提供原作结局。", GREEN)

    add_heading(doc, "页面结构", 2)
    add_table(doc,
        ["区域", "内容", "状态"],
        [
            ["阶段裁决头图", "案卷暂时封存 分数与阶段裁决", "固定显示版权和推演边界"],
            ["四卡证据链", "本局带入质证的四张证据", "按选中顺序或统一证据顺序显示"],
            ["假说选择", "世界规则置换 感知记忆错位 观察实验 证据不足", "选择后显示支持 缺口 下一步验证"],
            ["求证问题", "直接曝光 交叉验证 低风险试探", "显示信息价值 暴露风险 解释与局限"],
            ["底部动作", "查看原作 重新组牌 封存案卷完成", "完成按钮在两项都选择后启用"],
            ["刘看山", "解释推荐依据和安全提问", "强调推荐只来自牌组方向"],
        ], widths=[1.15, 3.4, 2.75], font_size=8.65, first_col_bold=True)

    add_heading(doc, "推荐逻辑", 2)
    add_table(doc,
        ["第四张方向牌", "默认推荐", "解释"],
        [
            ["城市资料", "世界规则发生置换", "牌组把身体异常与公共记录冲突连在一起"],
            ["灰夹克", "有人进行现实操控或观察实验", "牌组把定向测试与跟随行为连在一起"],
            ["未匹配", "证据不足 暂不站队", "程序找不到有效方向时使用保守结论"],
        ], widths=[1.6, 2.25, 3.45], font_size=8.75, first_col_bold=True)

    add_heading(doc, "求证问题", 2)
    bullet(doc, "直接曝光信息价值高、暴露风险高。")
    bullet(doc, "交叉验证是当前推荐项，兼顾独立记忆来源与暴露风险。")
    bullet(doc, "低风险试探暴露最少，但信息价值较低。")
    bullet(doc, "选择后提供复制问题按钮，复制成功提示约 1.6 秒。")

    add_heading(doc, "完成回执建议", 2)
    add_status(doc, "当前行为", "点击封存案卷完成后保存分数并返回入口页。入口页显示已完成和最佳裁决分数，但没有独立完成回执。", ORANGE)
    add_status(doc, "建议新增", "在返回入口前显示轻量完成回执。此项为 P1，不影响本轮 P0 演示。", BLUE)
    bullet(doc, "回执标题为 本局案卷已封存，显示分数、四卡证据链、所选假说、所选问题和来源模式。")
    bullet(doc, "主动作 返回故事页；次动作 再玩一次、复制求证问题、查看原作。")
    bullet(doc, "回执不得使用 通关真相 解锁结局 等措辞。")
    bullet(doc, "若云端保存失败，显示 已保存在本机，不阻断离开。")

    page_break(doc)

    add_heading(doc, "十四 加载 错误 空状态", 1)

    add_table(doc,
        ["状态", "当前实现", "建议呈现"],
        [
            ["故事列表加载", "居中加载文字", "P1 可替换为标题和主卡骨架，至少保留 300 px 高度"],
            ["故事详情加载", "加载文字与返回故事页", "保留返回按钮，不展示上一案卷残影"],
            ["列表或详情错误", "红色错误文字加重新加载", "错误说人话，按钮紧邻错误"],
            ["没有可玩故事", "今日互动案卷暂不可用", "提供重新加载，不展示空白地图"],
            ["证据库为空", "调查提示空状态", "提示去左侧点击场景热点"],
            ["质证回应失败", "对决区错误提示", "保留手牌与回合上下文，允许安全重试"],
            ["裁决失败", "结果区错误提示", "不改变本地胜负结果，允许重试请求"],
            ["资料同步失败", "非阻塞本地保存提示", "说明结果仍已保存在本机"],
        ], widths=[1.35, 2.45, 3.5], font_size=8.55, first_col_bold=True)

    add_heading(doc, "状态文案规则", 2)
    bullet(doc, "错误说明发生了什么、用户还能做什么，不显示堆栈和接口路径。")
    bullet(doc, "离线降级属于来源状态，不使用红色错误样式。")
    bullet(doc, "加载和提交按钮保留原尺寸，避免文字变化导致布局跳动。")
    bullet(doc, "所有异步反馈使用 role status 或 aria live；真正错误使用 role alert。")

    add_heading(doc, "数据与状态映射", 2)
    add_table(doc,
        ["数据或状态", "驱动界面", "空值或异常处理"],
        [
            ["storyFeed source", "入口在线或离线状态", "缺失时按离线处理并显示说明"],
            ["storyFeed warning", "透明降级原因", "没有 warning 时显示 Work ID"],
            ["featuredStory playable", "唯一可进入故事", "没有可玩项时显示重试"],
            ["demo source mode", "开场与来源阅读器模式", "只接受 zhihu live 或 curated fallback 的已知视觉"],
            ["verifiedQuoteCount", "短引核验徽标", "零时不得显示已核验"],
            ["discovered", "热点 证据库 搜证 KPI", "可逆操作必须同步清除已选状态"],
            ["selectedEvidence", "四卡牌组 对决初始手牌 End 证据链", "最多四张"],
            ["keyEvidenceIds", "必选锚点", "缺失时进入质证失败并列出名称"],
            ["hypothesisEvidenceIds", "第四张方向牌范围", "背景卡禁用"],
            ["battle", "生命 体力 护盾 回合 手牌 胜负", "任何阶段切换前后保持一致"],
            ["verdict", "阶段裁决和分数", "请求失败不改变 battleResult"],
            ["storyEnding", "四假说 三问题 End 文案", "缺失时不进入 End"],
        ], widths=[1.55, 3.0, 2.75], font_size=8.25, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "十五 桌面与移动端适配", 1)
    paragraph(doc, "当前 CSS 已覆盖 1100、900、820、680、560 和 420 px 等断点。UI 交付可归并为四档设计稿，前端继续保留细分规则。")
    add_table(doc,
        ["设计档", "建议宽度", "主要布局"],
        [
            ["宽桌面", "1440 px", "入口双列 调查三栏 对决四卡横排"],
            ["桌面或平板横屏", "1024 px", "适当压缩栏宽 保持主要信息同屏"],
            ["移动端", "390 px", "调查上方双列 下方证据牌组 刘看山面板近全宽"],
            ["紧凑移动端", "360 px", "标题压缩 对决两列卡片 开场和 End 全部单列"],
        ], widths=[1.45, 1.35, 4.5], font_size=8.8, first_col_bold=True)

    add_heading(doc, "移动端规则", 2)
    bullet(doc, "搜证页在 680 px 以下改为两列上区和一列下区：故事现场、来源阅读器并列，证据牌组横跨整行。")
    bullet(doc, "对决在 800 px 以下手牌改为两列；560 px 以下中间对话区移到角色区上方。")
    bullet(doc, "故事开场在 560 px 以下封面改为横向短图，按钮上下排列。")
    bullet(doc, "End 在 820 px 以下头图、假说和问题卡改为单列或两列自适应。")
    bullet(doc, "刘看山首页卡在 560 px 以下上下堆叠；悬浮触发器只保留角色头像。")
    bullet(doc, "固定按钮使用 env safe area inset，避免被底部手势条遮挡。")

    add_heading(doc, "实机验收矩阵", 2)
    add_table(doc,
        ["设备场景", "必须验证"],
        [
            ["1440 乘 900", "入口首屏 主 CTA 调查三栏 对决四卡和 End 完整可用"],
            ["1024 乘 768", "三栏不出现横向滚动，来源阅读器和牌组仍可操作"],
            ["390 乘 844", "刘看山面板不遮按钮，搜证和牌组滚动区域清晰"],
            ["360 乘 800", "标题不截断，卡牌文字可读，固定按钮不覆盖底部操作"],
            ["横屏手机", "高度较小时模态框内部滚动，对决上下文不丢失"],
            ["浏览器缩放 200 百分比", "主要动作和来源状态仍可访问，无水平滚动"],
        ], widths=[1.8, 5.5], font_size=8.75, first_col_bold=True)

    add_status(doc, "待完成", "响应式 CSS 已存在，但小屏触控目标、双滚动和移动浏览器地址栏仍需真机验收。", ORANGE)

    page_break(doc)

    add_heading(doc, "十六 刘看山素材与动效映射", 1)
    add_figure(doc, ASSETS / "turnaround.jpg", 6.75, "刘看山三视图参考", "刘看山正面 侧面 背面的三维造型参考图")

    add_table(doc,
        ["素材", "当前使用位置", "说明"],
        [
            ["greeting gif", "首页默认导览和悬浮入口", "第一次见面和未展开状态"],
            ["idle gif", "悬浮向导展开后的触发器", "保持陪伴感，不抢正文"],
            ["researching gif", "首页任务与边界主题 对决提示", "表达查资料与梳理论证"],
            ["thinking gif", "搜证与 End 面板", "表达比较证据和选择下一步"],
            ["celebrate gif", "对决胜利与 End 触发器", "只在完成节点使用"],
            ["sleeping gif", "当前未使用", "P2 可用于长时间无操作，不进入本轮 P0"],
            ["still png", "减少动态效果", "所有 GIF 的静态替代"],
            ["portrait jpg", "设计参考与文档封面", "当前网页未直接使用"],
            ["turnaround jpg", "造型参考", "当前网页未直接使用"],
        ], widths=[1.45, 2.65, 3.2], font_size=8.55, first_col_bold=True)

    add_heading(doc, "动效规则", 2)
    bullet(doc, "同一角色区域一次只播放一个 GIF，不叠加多个动作。")
    bullet(doc, "角色动画不承载唯一信息；标题、正文和状态文字必须独立完整。")
    bullet(doc, "系统启用 prefers reduced motion 时隐藏 GIF 并显示 still png，同时取消 hover 位移。")
    bullet(doc, "动画文件约 320 乘 320 且 20 fps，应懒加载非首屏素材，避免同时解码。")
    bullet(doc, "角色始终保持白色主体与干净背景，不用滤镜改变知乎 IP 造型。")

    page_break(doc)

    add_heading(doc, "十七 无障碍与可用性", 1)

    add_heading(doc, "当前已有", 2)
    bullet(doc, "跳到主要内容链接、页面语义标签和清晰的焦点外框。")
    bullet(doc, "语言切换、主题标签、证据选择和假说选择使用 aria selected 或 aria pressed。")
    bullet(doc, "刘看山答案、问题结果和异步状态使用 aria live。")
    bullet(doc, "错误使用 role alert，加载使用 role status。")
    bullet(doc, "刘看山动画支持 prefers reduced motion 静态替代。")

    add_heading(doc, "P0 必须补验", 2)
    bullet(doc, "故事开场打开后焦点进入对话框，Tab 不离开，关闭后回到触发按钮。", checkbox=True)
    bullet(doc, "所有图标按钮有中文和英文 aria label。", checkbox=True)
    bullet(doc, "来源在线、离线、已核验和策划转述不仅依赖颜色。", checkbox=True)
    bullet(doc, "证据卡、热点和刘看山问题按钮的触控区域至少 44 乘 44 px。", checkbox=True)
    bullet(doc, "正文与背景对比不低于 4.5 比 1，大字号不低于 3 比 1。", checkbox=True)
    bullet(doc, "键盘可完成从入口到 End 的全部流程。", checkbox=True)
    bullet(doc, "对决倒计时不会以每秒朗读打断屏幕阅读器；只在关键阈值提示。", checkbox=True)
    bullet(doc, "音频提供可见静音入口，并遵守浏览器自动播放限制。", checkbox=True)

    add_heading(doc, "文案可用性", 2)
    bullet(doc, "按钮写结果，例如 开始搜证 返回调查重组 复制问题，不只写 确定。")
    bullet(doc, "错误包含缺少的数量或证据名称。")
    bullet(doc, "专业概念在首次出现时解释，后续统一使用 关键锚点 方向牌 阶段假说。")
    bullet(doc, "中英文切换后不混用固定中文来源标签；品牌名和 Work ID 可保留。")

    add_heading(doc, "建议测试", 2)
    add_status(doc, "建议新增", "把键盘路径、减少动态效果、200 百分比缩放和屏幕阅读器快速检查加入演示前检查表。", BLUE)
    add_status(doc, "建议新增", "用自动化工具检查可访问名称、对比度和焦点顺序，但最终以人工完成主流程为准。", BLUE)

    page_break(doc)

    add_heading(doc, "十八 实现优先级", 1)

    add_heading(doc, "P0 黑客松交付必须完成", 2)
    bullet(doc, "统一入口、开场、搜证、组牌、对决、失败和 End 的字体、边框、按钮与状态样式。")
    bullet(doc, "保证在线短引、在线未命中、离线转述三种来源模式准确显示。")
    bullet(doc, "保证六条线索、三张关键锚点、一张方向牌和四卡出示门槛无视觉歧义。")
    bullet(doc, "完成 1440、1024、390 和 360 宽度验收，消除遮挡、溢出和双滚动。")
    bullet(doc, "完成故事开场的焦点管理、键盘主流程和减少动态效果验收。")
    bullet(doc, "确认失败可恢复、End 按钮门槛、完成后分数保存和入口最佳分数更新。")

    add_heading(doc, "P1 演示后优先补齐", 2)
    bullet(doc, "增加完成回执，保留本局假说、问题、四卡和来源模式。")
    bullet(doc, "增加音频静音开关并记住用户选择。")
    bullet(doc, "加载文字升级为稳定骨架，减少布局跳动。")
    bullet(doc, "失败分析高亮未出示证据并提供定位重组。")
    bullet(doc, "增加关键漏斗埋点：进入调查、搜齐、组牌完成、对决胜负、End 完成。")

    add_heading(doc, "P2 后续探索", 2)
    bullet(doc, "扩展更多故事时复用数据驱动场景和来源标识，不复制页面。")
    bullet(doc, "若接入自由对话 Agent，新增引用、拒答、会话失败与未核验提示，不替换当前稳定预设。")
    bullet(doc, "考虑睡眠动作、故事进度地图或可分享回执，但不能增加 P0 演示路径长度。")

    add_heading(doc, "交付顺序", 2)
    add_table(doc,
        ["顺序", "交付物", "通过标准"],
        [
            ["一", "视觉 token 和通用组件", "入口、调查、对决、End 使用同一套状态"],
            ["二", "入口和刘看山", "一分钟内理解故事与玩法"],
            ["三", "搜证与来源阅读器", "六线索和来源模式无歧义"],
            ["四", "组牌与对决", "门槛、资源和胜负清楚"],
            ["五", "End 与失败", "推演边界、恢复和完成可用"],
            ["六", "移动端与无障碍", "主流程在目标设备完整通过"],
        ], widths=[0.72, 2.45, 4.13], font_size=8.8, first_col_bold=True)

    page_break(doc)

    add_heading(doc, "十九 UI 验收清单", 1)
    paragraph(doc, "以下项目全部通过后，才把界面标记为可演示。建议由 UI 设计、前端和内容负责人共同走查一次在线模式和一次离线模式。")

    add_heading(doc, "内容与来源", 2)
    for item in [
        "入口、开场、来源阅读器和 End 都显示故事名、作者和内容边界。",
        "在线命中短引与策划转述在同一页可被清楚区分。",
        "离线模式不出现 已核验原文 或 暗示接口在线。",
        "页面没有 正确答案 真相揭晓 解锁原作结局 等误导文案。",
        "查看原作使用真实外链且在新窗口打开。",
    ]:
        bullet(doc, item, checkbox=True)

    add_heading(doc, "入口与刘看山", 2)
    for item in [
        "入口首屏在目标桌面尺寸能看到故事钩子、来源状态和主按钮。",
        "刘看山三主题切换不引起明显布局跳动。",
        "首页导览和悬浮向导均能用键盘操作。",
        "调查、对决、胜利、失败和 End 的向导内容与当前状态一致。",
        "减少动态效果模式只显示静态素材。",
    ]:
        bullet(doc, item, checkbox=True)

    add_heading(doc, "搜证与组牌", 2)
    for item in [
        "三场景共有六条且不会重复计数。",
        "撤销线索会同步移除证据库和牌组选择。",
        "培训教材明确为背景线索并不可入庭。",
        "四卡牌组必须包含三张关键锚点和一张方向牌。",
        "所有规则错误说明缺少数量或具体证据。",
    ]:
        bullet(doc, item, checkbox=True)

    add_heading(doc, "对决与 End", 2)
    for item in [
        "倒计时、论证稳定度、专注、缓冲和日常解释力都同时显示文本与数值。",
        "四张证据未全部出示时，重复强牌不能结束对决。",
        "失败后可返回调查且搜证进度不丢失。",
        "End 推荐假说与第四张方向牌一致，并明确只是阶段推演。",
        "未选择假说或问题时完成按钮保持禁用。",
        "完成后分数正确保存并在入口显示最佳成绩。",
    ]:
        bullet(doc, item, checkbox=True)

    add_heading(doc, "响应式与稳定性", 2)
    for item in [
        "1440、1024、390 和 360 宽度无横向滚动、遮挡或裁切。",
        "刘看山悬浮面板不遮挡主要操作，移动端内部滚动可用。",
        "所有点击目标满足触控尺寸，键盘焦点始终可见。",
        "在线接口超时后自动进入透明离线模式，完整主流程仍可完成。",
        "图片加载失败有稳定占位，不造成布局崩塌。",
        "背景音乐、反馈音与退出清理正常，不发生叠播。",
    ]:
        bullet(doc, item, checkbox=True)

    add_heading(doc, "最终交付定义", 2)
    paragraph(doc, "P0 清单全部通过，在线和离线两条来源路径均完成一次从入口到 End 的人工走查，且 UI 同学已经把任何建议新增项明确标为后续，才可将《蓝血》互动案卷认定为本次黑客松的完整可交付产品。")

    # Core properties and table cell normalization
    props = doc.core_properties
    props.title = "蓝血互动游戏 UI 实现说明"
    props.subject = "知乎黑客松蓝血互动案卷的 UI 与前端实现规范"
    props.author = "蓝血互动案卷项目组"
    props.keywords = "蓝血, 刘看山, 知乎黑客松, UI, 交互游戏"
    props.comments = "区分当前实现与建议后续"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    set_paragraph_keep(p)

    doc.save(str(OUTPUT))


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
