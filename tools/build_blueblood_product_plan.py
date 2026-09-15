from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import build_blueblood_ui_spec as ui  # noqa: E402


OUTPUT = ROOT / "docs" / "蓝血互动叙事游戏产品说明计划书.docx"
ASSETS = ROOT / "frontend" / "public" / "assets" / "kanshan"


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("蓝血互动叙事游戏产品说明计划书   ")
    ui.set_run_font(run, size=8.3, color="6D685F")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ui.BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), ui.BODY_FONT)
    r_pr.append(r_fonts)
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = ui.BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), ui.BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), ui.BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), ui.BODY_FONT)
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor.from_string(ui.INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = ui.BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), ui.BODY_FONT)
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(ui.INK)
    title.paragraph_format.space_after = Pt(13)
    title_p_pr = title._element.get_or_add_pPr()
    old_border = title_p_pr.find(qn("w:pBdr"))
    if old_border is not None:
        title_p_pr.remove(old_border)
    border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        border.append(item)
    title_p_pr.append(border)

    for level, size in ((1, 18), (2, 13.5), (3, 11.5)):
        style = styles[f"Heading {level}"]
        style.font.name = ui.BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), ui.BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ui.INK)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6 if level == 1 else 4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name, size, after in (("Plan Bullet", 10.1, 3.5), ("Plan Checklist", 10, 4), ("Plan Status", 9.8, 4)):
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = ui.BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), ui.BODY_FONT)
        style.font.size = Pt(size)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.22

    add_page_number(section.footer.paragraphs[0])
    return doc


def paragraph(doc, text: str = "", *, size=None, color=None, bold_lead=None, align=None, space_after=5, line=1.25):
    return ui.paragraph(
        doc,
        text,
        bold_lead=bold_lead,
        color=color,
        size=size,
        align=align,
        space_after=space_after,
        line=line,
    )


def bullet(doc, text: str, *, checkbox=False, color=None):
    p = doc.add_paragraph(style="Plan Checklist" if checkbox else "Plan Bullet")
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(("□ " if checkbox else "• ") + text)
    ui.set_run_font(r, size=10.1, color=color or ui.INK)
    ui.set_paragraph_keep(p)
    return p


def heading(doc, text: str, level=1, page_break=False):
    return ui.add_heading(doc, text, level, page_break=page_break)


def status(doc, label: str, text: str, color: str):
    p = doc.add_paragraph(style="Plan Status")
    r = p.add_run(label)
    ui.set_run_font(r, size=9.2, bold=True, color=color)
    r = p.add_run("  " + text)
    ui.set_run_font(r, size=9.8, color=ui.INK)
    ui.set_paragraph_keep(p)


def table(doc, headers, rows, widths, *, font_size=8.6, first_col_bold=False):
    return ui.add_table(
        doc,
        headers,
        rows,
        widths=widths,
        header_fill=ui.DARK_BLUE,
        font_size=font_size,
        first_col_bold=first_col_bold,
    )


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(38)
    r = p.add_run("ZHIHU HACKATHON   PRODUCT PLAN")
    ui.set_run_font(r, name=ui.MONO_FONT, size=9.3, bold=True, color=ui.BLUE)

    p = doc.add_paragraph(style="Title")
    r = p.add_run("蓝血互动叙事游戏产品说明计划书")
    ui.set_run_font(r, size=29, bold=True, color=ui.INK)
    p.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("面向知乎黑客松四人团队的产品范围 实施计划与验收标准")
    ui.set_run_font(r, size=13.5, bold=True, color=ui.GRAY)

    ui.add_figure(
        doc,
        ASSETS / "portrait.jpg",
        1.72,
        "刘看山故事向导形象",
        "白色背景中的刘看山三维角色立姿，用于蓝血互动叙事游戏的故事导览",
    )

    meta = [
        ("产品", "蓝血疑云 单故事互动案卷"),
        ("版本", "1.0"),
        ("日期", "2026 年 9 月 15 日"),
        ("计划周期", "四人并行 四小时完成黑客松交付"),
        ("演示入口", "https://zhihucat-six.vercel.app/campaign"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(label + "  ")
        ui.set_run_font(r, size=9.4, bold=True, color=ui.BLUE)
        if label == "演示入口":
            add_hyperlink(p, "打开线上演示", value)
        else:
            r = p.add_run(value)
            ui.set_run_font(r, size=10.2, color=ui.INK)

    paragraph(
        doc,
        "交付结论  本轮只维护《蓝血》一个故事，先保证入口到 End 的三分钟闭环、来源可追溯和离线可演示，再处理非关键扩展。最终完成状态必须以合并后的构建、测试和人工走查结果为准。",
        size=10.5,
        color=ui.GRAY,
        space_after=0,
        line=1.42,
    )


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = configure_document()
    add_cover(doc)
    doc.add_page_break()

    heading(doc, "一 执行摘要", 1)
    paragraph(doc, "《蓝血疑云》把知乎故事《蓝血》的公开内容线索改编为一局可质证的互动案卷。玩家先听刘看山说明故事和内容边界，再进入三处场景寻找六条线索，选择四张证据牌挑战日常解释，最后选择一个阶段假说和一个知乎求证问题。目标体验时长约三分钟。")
    paragraph(doc, "产品不承诺复现原作全文或揭示原作结局。在线接口只用于读取故事目录、详情和核验预设短引；接口不可用或短引未命中时，页面明确切换为策划转述，并继续完成同一条玩法流程。", bold_lead="产品不承诺")

    heading(doc, "当前决策", 2)
    table(doc, ["决策", "本轮做法", "原因"], [
        ["故事范围", "只维护《蓝血》一个可玩故事", "四小时内保证闭环、来源和演示稳定"],
        ["刘看山形态", "确定性故事向导，不开放自由聊天", "降低幻觉、延迟和临场不可控风险"],
        ["内容使用", "仅展示接口即时核验的短引，其余为策划转述", "尊重内容边界并保持来源可信"],
        ["结局设计", "阶段推演加求证问题，不宣称原作答案", "让玩家带着问题回到知乎，而非替作者续写"],
        ["交付优先级", "P0 完整主流程优先，P1 视觉补强，P2 后续探索", "减少四小时冲刺中的范围漂移"],
    ], [1.25, 2.65, 3.15], font_size=8.75, first_col_bold=True)

    heading(doc, "当前状态", 2)
    status(doc, "代码路径已存在", "入口、刘看山导览、搜证、组牌、对决、失败和 End 的主要界面与数据结构已经存在，但最终交付仍须在合并后的分支重新构建和走查。", ui.GREEN)
    status(doc, "已接入待联调", "知乎故事列表与详情客户端、蓝血案卷转换和透明降级代码已经存在；线上环境、代理路由和最终部署版本仍需确认。", ui.ORANGE)
    status(doc, "尚未完成", "完整回执页、多故事案卷和自由对话 Agent 不属于当前可交付能力。", ui.RED)

    heading(doc, "文档导航", 2)
    paragraph(
        doc,
        "二至四讲产品、创新和内容边界；五至七讲玩法、刘看山 Agent 与全部页面；八至九讲知乎 API、降级和技术边界；十讲四人四小时分工；十一至十二讲验收与风险；十三至十四提供演示讲稿和提交清单。",
        size=9.7,
        color=ui.GRAY,
        line=1.32,
    )

    doc.add_page_break()
    heading(doc, "二 产品定位与目标用户", 1)
    heading(doc, "产品定位", 2)
    paragraph(doc, "产品定位为知乎故事的互动案卷体验。它把阅读材料转化为有门槛、有反馈、可回退的推理流程，同时让每条材料保留清楚的来源状态。演示重点是玩家如何理解、选择和质证，而不是展示一个自动生成结局的聊天机器人。")

    heading(doc, "目标用户", 2)
    table(doc, ["用户", "当前痛点", "产品回应"], [
        ["悬疑故事读者", "故事信息多但参与方式通常只有阅读和评论", "用搜证、组牌和质证把阅读转成可操作任务"],
        ["轻度游戏用户", "复杂规则学习成本高，短时体验难进入", "刘看山先解释故事，再用六条线索和四张牌控制规则规模"],
        ["知乎内容探索者", "容易混淆原文、转述和二次推演", "每个材料显示来源模式、作者、位置和核验状态"],
        ["黑客松评委", "需要快速看到 API 价值、产品闭环和稳定性", "三分钟完成一局，在线失败时透明降级而不阻断演示"],
    ], [1.3, 2.7, 3.1], font_size=8.7, first_col_bold=True)

    heading(doc, "核心价值", 2)
    bullet(doc, "理解更快：入口和刘看山在一分钟内说明世界观、任务和不能确定的部分。")
    bullet(doc, "参与更深：玩家必须发现、选择并实际出示证据，不能直接跳到结论。")
    bullet(doc, "来源更清楚：在线短引、未命中材料和离线转述使用不同文字标识。")
    bullet(doc, "演示更稳定：接口异常时保留完整玩法，并向玩家说明降级原因。")

    heading(doc, "目标与非目标", 2)
    table(doc, ["范围", "本轮目标", "本轮不做"], [
        ["内容", "完成《蓝血》单故事闭环", "抓取或保存原作完整正文"],
        ["交互", "确定性引导、搜证、组牌、对决和阶段推演", "开放式角色扮演和无限分支剧情"],
        ["数据", "使用知乎公开接口和可追溯的策划案卷", "把推演包装成作者原意或唯一真相"],
        ["平台", "桌面和移动网页可演示", "原生客户端、多人联机和内容后台"],
    ], [1.0, 3.15, 2.95], font_size=8.8, first_col_bold=True)

    heading(doc, "三 核心创新", 1)
    table(doc, ["创新", "具体设计", "对比赛的可见价值"], [
        ["刘看山叙事向导", "用三组固定主题讲清故事、玩法和不确定性；在搜证、对决和 End 中提供上下文提示", "把知乎 IP 角色变成可用的产品导航，并避免自由对话幻觉"],
        ["可溯源证据玩法", "每张证据同时包含证明目的、可信度、来源位置和核验状态", "让内容 API 参与核心玩法，而不是只做故事列表"],
        ["六选四质证", "三张关键锚点固定，一张方向牌由玩家选择；对决要求四张都出示", "在短局内制造真实取舍，避免只刷最高数值"],
        ["阶段推演 End", "玩家选择竞争假说，再选择不同信息价值和暴露风险的知乎问题", "结局从猜答案变为设计下一步验证"],
        ["透明降级", "在线命中、在线未命中和离线案卷三种状态共用主流程", "网络或接口异常时仍能演示，同时不冒充在线来源"],
    ], [1.35, 3.45, 2.3], font_size=8.55, first_col_bold=True)

    heading(doc, "创新成立的判断标准", 2)
    bullet(doc, "评委能指出知乎内容接口具体影响了哪些材料和状态，而不是只看到品牌标识。")
    bullet(doc, "玩家能解释为什么第四张方向牌会改变阶段推演，但不会改变原作事实。")
    bullet(doc, "断网演示仍可完成，且页面不会显示已核验原文或假装接口在线。")
    bullet(doc, "刘看山帮助玩家理解下一步，不替玩家做选择，也不回答超出案卷的事实问题。")

    doc.add_page_break()
    heading(doc, "四 故事框架与内容边界", 1)
    heading(doc, "故事概况", 2)
    table(doc, ["项目", "内容"], [
        ["故事", "《蓝血》"],
        ["作者", "桃花先生"],
        ["知乎 Work ID", "2025684191967294692"],
        ["开场钩子", "方诺发现周围所有人都把血液原本是蓝色当作常识，而她自己的血从一开始就是红色"],
        ["玩家身份", "方诺的匿名调查协助者"],
        ["玩家目标", "比较客观异常与疲劳、误记、巧合两类解释，形成阶段证据链并决定下一步求证"],
    ], [1.4, 5.65], font_size=8.85, first_col_bold=True)
    paragraph(doc, "故事简介  方诺先在急救培训中发现血液常识冲突，随后发现同事的血由蓝转红、公共地理资料与自己的记忆不一致、培训方给她单独设置常识测试，并遇到疑似尾随者。游戏只把这些已公开线索组织成可玩案卷，不补写异常机制和后续结局。", bold_lead="故事简介")

    heading(doc, "六条案卷线索", 2)
    table(doc, ["场景", "线索", "功能", "能否上庭"], [
        ["急救培训室", "培训教材的蓝血常识", "建立共同背景", "否"],
        ["急救培训室", "指尖流出的红血", "关键身体异常", "必选"],
        ["急救培训室", "同事由蓝变红的血迹", "独立观察锚点", "必选"],
        ["深夜资料库", "与记忆冲突的城市资料", "世界规则方向", "二选一"],
        ["深夜资料库", "只针对方诺的测试卷", "主动筛查锚点", "必选"],
        ["老街死胡同", "没有折返的灰夹克", "观察实验方向", "二选一"],
    ], [1.35, 2.45, 2.1, 1.15], font_size=8.7, first_col_bold=True)

    heading(doc, "内容层级", 2)
    table(doc, ["层级", "允许表达", "禁止表达"], [
        ["原文事实", "接口即时命中的公开短引，并显示作者、位置和核验状态", "把预设转述或缓存内容称为逐字原文"],
        ["策划转述", "根据公开简介和已知世界观整理的线索说明", "隐藏转述身份或暗示内容来自实时接口"],
        ["竞争假说", "世界置换、感知或记忆错位、受控观察、证据不足", "把任何假说写成正确答案"],
        ["阶段裁决", "本局证据链更支持哪种解释及仍缺什么", "宣称解锁原作结局或替作者续写"],
    ], [1.15, 3.2, 2.7], font_size=8.65, first_col_bold=True)

    heading(doc, "版权与数据边界", 2)
    bullet(doc, "故事与角色归原作者和相应权利人；页面持续显示作者、来源和改编说明。")
    bullet(doc, "完整正文不落盘。代码只在实时详情中核验预设短引是否出现，并返回命中结果。")
    bullet(doc, "原作入口使用游戏内查看原作按钮，打开知乎页面；阶段推演不会替代原作阅读。")
    bullet(doc, "任何新增文案都必须标明是原文短引、策划转述还是游戏推演。")

    doc.add_page_break()
    heading(doc, "五 完整玩法与用户旅程", 1)
    heading(doc, "三分钟主流程", 2)
    paragraph(doc, "入口页 → 刘看山导览 → 故事开场 → 三场景搜证 → 六选四组牌 → 解释力对决 → 阶段推演 End → 封存并返回")
    table(doc, ["阶段", "玩家任务", "系统反馈", "离开条件"], [
        ["入口", "理解故事钩子、来源状态和内容边界", "显示 3 场景 6 线索 4 卡牌 1 求证", "进入刘看山导览"],
        ["刘看山导览", "查看故事、玩法和不确定性三项说明", "切换主题与角色动作", "点击开始搜证"],
        ["故事开场", "确认身份、任务和来源", "模态卡显示简介与版权说明", "点击开始搜证"],
        ["搜证", "在三个场景找齐六条线索", "热点、证据库和进度同步更新", "发现六条线索"],
        ["组牌", "选择三张锚点和一张方向牌", "显示 4 比 4 和缺失规则", "四张牌满足门槛"],
        ["对决", "用四张证据削弱日常解释并管理资源", "回合、体力、专注、护盾和对话更新", "解释力归零且四张均出示"],
        ["失败", "查看链条缺口并返回重组", "保留搜证和牌组进度", "回到调查"],
        ["End", "选择阶段假说和知乎求证问题", "显示依据、缺口、风险和下一步", "两项都选择后封存"],
    ], [0.9, 2.25, 2.55, 1.35], font_size=8.35, first_col_bold=True)

    heading(doc, "关键玩法规则", 2)
    bullet(doc, "必须先发现六条线索，才能带四张牌进入质证。")
    bullet(doc, "三张关键锚点必须入选，第四张从城市资料和灰夹克中二选一。")
    bullet(doc, "培训教材属于背景线索，不能作为第四张牌。")
    bullet(doc, "对决中四张入选证据都必须至少出示一次，重复使用强牌不能直接结束。")
    bullet(doc, "End 必须同时选择阶段假说和求证问题；未选择时完成按钮保持禁用。")

    heading(doc, "核心资源与反馈", 2)
    table(doc, ["资源", "作用", "失败时如何恢复"], [
        ["日常解释力", "对手生命值，归零后进入阶段推演", "重组牌组后重新开始"],
        ["玩家体力", "承受反方质疑，归零则本轮失败", "使用恢复牌或重新开始"],
        ["专注", "支付出牌成本，满值时仍需正常出牌", "恢复牌增加专注"],
        ["反驳缓冲", "抵消下一次部分伤害", "防御牌增加缓冲"],
        ["回合倒计时", "限制单回合思考时间", "超时进入反方行动，但不清空进度"],
    ], [1.2, 3.3, 2.55], font_size=8.75, first_col_bold=True)

    heading(doc, "End 的四种阶段假说", 2)
    table(doc, ["假说", "当前支持", "必须保留的缺口"], [
        ["世界规则发生置换", "两种血液现象和公共资料冲突", "没有机制、发生时间和独立经历者"],
        ["感知或记忆系统性错位", "周围资料与他人认知彼此一致", "无法完整解释身体观察和定向试卷"],
        ["受控观察或现实操控", "定向测试与持续跟随", "两者之间没有直接因果证据"],
        ["证据不足暂不站队", "所有解释只覆盖部分线索", "仍需独立记录和第三方验证"],
    ], [1.6, 2.75, 2.7], font_size=8.65, first_col_bold=True)

    doc.add_page_break()
    heading(doc, "六 刘看山故事 Agent", 1)
    paragraph(doc, "刘看山承担故事向导、规则解释和来源提醒三项职责。当前版本使用确定性预设文案，不调用自由生成模型。这样可以保证三分钟演示中的回答一致，并防止角色把推演说成原作事实。")

    heading(doc, "入口导览", 2)
    table(doc, ["主题", "必须回答的问题", "结束动作"], [
        ["故事发生了什么", "蓝血常识、方诺红血、同事蓝转红、资料冲突、定向测试和跟随者", "继续切换或开始搜证"],
        ["我在游戏里做什么", "三场景、六线索、四张牌、对决、阶段假说和知乎问题", "开始搜证"],
        ["哪些还不能确定", "多种假说都只是阶段推演，原作结局未在本产品中给出", "开始搜证"],
    ], [1.55, 4.45, 1.05], font_size=8.7, first_col_bold=True)

    heading(doc, "上下文向导", 2)
    table(doc, ["所在页面", "默认帮助", "不得做的事"], [
        ["搜证", "解释事实与推断、证据强度和第四张牌方向", "替玩家自动收集或直接给牌组答案"],
        ["对决", "说明反方论点、重复出牌和资源变化", "承诺某个假说是真相"],
        ["失败", "指出缺少的证据和返回路径", "把玩法失败描述为假说错误"],
        ["End", "比较推荐依据、证据缺口和提问风险", "替作者补写后续或解锁结局"],
        ["离线状态", "说明当前材料是策划转述", "显示已核验短引标识"],
    ], [1.15, 3.75, 2.15], font_size=8.7, first_col_bold=True)

    heading(doc, "未来接入 Astra 的边界", 2)
    paragraph(
        doc,
        "后续若用 Astra 做自由问答，输入只包含当前案卷、已发现线索、来源状态和玩家阶段；事实性输出必须附材料 ID，并区分事实、推断和未知。无材料、超时或来源校验失败时回退到当前确定性答案，且不得写回原始故事数据。本轮不把自由问答列入完成功能或演示必经路径。",
        size=9.8,
        line=1.28,
    )

    heading(doc, "七 页面与 UI 说明", 1)
    table(doc, ["页面", "核心信息", "主要动作", "关键状态"], [
        ["全局顶栏", "品牌、语言、玩家和积分", "返回入口、切换语言、账户操作", "加载、登录、同步失败"],
        ["故事入口", "钩子、来源、玩法数字、内容边界", "听刘看山讲故事、重新调查", "在线、离线、已完成"],
        ["刘看山导览", "三项主题和开始按钮", "切换主题、开始搜证", "默认、切换、减少动态效果"],
        ["故事开场", "身份、任务、简介、版权", "返回、开始搜证", "实时接口、离线降级"],
        ["调查页", "三场景、来源阅读器、证据库", "找线索、看来源、选牌", "未发现、已发现、已选、背景锁定"],
        ["质证页", "对手、对话、资源、手牌", "出牌、恢复、防御", "玩家行动、反方行动、胜负"],
        ["失败页", "证据链缺口和本轮分析", "返回调查重组", "玩法失败、接口错误"],
        ["End", "四卡链、四假说、三问题", "选择、复制问题、查看原作、完成", "未选择、可完成、已封存"],
    ], [1.0, 2.35, 2.2, 1.5], font_size=8.3, first_col_bold=True)

    heading(doc, "统一视觉语言", 2)
    table(doc, ["项目", "基线", "验收重点"], [
        ["视觉风格", "纸张底、黑色粗边框、蓝黄强调、硬阴影", "保持案卷感，不让装饰压过来源与任务"],
        ["信息层级", "展示标题、页面标题、组件标题、正文、状态标签", "来源状态和下一步动作始终高于装饰元素"],
        ["交互反馈", "文字、轮廓和颜色共同表达", "不能只依赖颜色区分选择、错误和来源"],
        ["动效", "150 至 250 ms，角色区域一次一个 GIF", "减少动态效果时使用 still png"],
        ["移动端", "390 和 360 宽度单列或上下分区", "不横向滚动，不遮挡固定按钮"],
    ], [1.15, 3.05, 2.85], font_size=8.65, first_col_bold=True)

    heading(doc, "刘看山素材使用", 2, page_break=True)
    table(doc, ["素材", "推荐位置", "说明"], [
        ["greeting gif", "入口默认导览", "第一次见面"],
        ["idle gif", "悬浮向导展开后", "陪伴状态"],
        ["researching gif", "搜证与资料主题", "查资料和梳理论证"],
        ["thinking gif", "组牌和 End", "比较证据和下一步"],
        ["celebrate gif", "对决胜利", "只在完成节点使用"],
        ["still png", "减少动态效果模式", "所有 GIF 的静态替代"],
    ], [1.45, 2.3, 3.3], font_size=8.75, first_col_bold=True)

    doc.add_page_break()
    heading(doc, "八 知乎内容 API 与降级架构", 1)
    heading(doc, "调用链", 2)
    paragraph(doc, "玩家浏览器 → Next.js 页面 → 同源代理 /argus-api → Node API → 知乎黑客松故事 API")
    paragraph(doc, "异常分支 → Node API 返回 curated fallback 蓝血案卷 → 前端显示策划离线模式 → 玩家继续完成同一条主流程", color=ui.GRAY)

    heading(doc, "接口范围", 2)
    table(doc, ["层级", "接口或配置", "用途", "当前边界"], [
        ["知乎上游", "GET /story/list", "读取可见故事目录", "只有《蓝血》被编排为可玩案卷"],
        ["知乎上游", "GET /story/{workId}", "读取标题、作者、简介、标签和正文响应", "正文只在内存中核验预设短引，不落盘"],
        ["项目后端", "GET /api/zhihu/stories", "返回故事选择和来源状态", "接口失败时只返回《蓝血》离线条目"],
        ["项目后端", "GET /api/zhihu/stories/:id/case", "返回蓝血案卷、线索、假说和求证问题", "其他 Work ID 返回尚未编排"],
        ["前端", "NEXT_PUBLIC_API_BASE_URL", "配置同源代理或后端地址", "生产环境需与代理规则一致"],
    ], [1.1, 2.15, 2.45, 1.35], font_size=8.25, first_col_bold=True)

    heading(doc, "输入与输出字段", 2)
    table(doc, ["对象", "关键字段", "使用方式"], [
        ["故事摘要", "work_id title artwork description labels", "入口列表和故事卡"],
        ["故事详情", "chapter_name author_name introduction content", "构建来源信息并核验短引"],
        ["案卷来源", "mode verifiedQuoteCount notice originalUrl", "驱动在线、未命中和离线 UI"],
        ["案卷内容", "scenes documents evidence keyEvidenceIds", "驱动搜证、阅读器和组牌"],
        ["阶段推演", "hypotheses questions judgment", "驱动 End 和阶段裁决"],
    ], [1.25, 3.2, 2.6], font_size=8.65, first_col_bold=True)

    heading(doc, "三种来源状态", 2)
    table(doc, ["状态", "判定", "页面文案", "允许展示"], [
        ["在线且命中", "实时详情包含一个或多个预设短引", "知乎 API 已连接 已核验 N 条短引", "命中短引、位置、作者和策划释义"],
        ["在线但零命中", "详情可读但预设短引未匹配", "原文锚点未命中 策划线索转述", "转述和未命中说明"],
        ["离线案卷", "列表或详情失败、超时或格式异常", "策划离线模式 透明降级", "内置案卷和恢复说明"],
    ], [1.25, 2.15, 2.3, 1.35], font_size=8.4, first_col_bold=True)

    heading(doc, "降级原则", 2)
    bullet(doc, "当前内容客户端的默认上游超时时间为八秒；超时或无效响应进入可解释的降级路径。")
    bullet(doc, "降级不会阻断入口、搜证、组牌、对决和 End，但必须清除已核验短引标识。")
    bullet(doc, "verifiedQuoteCount 只能由后端命中结果产生，前端不得根据文案猜测。")
    bullet(doc, "远程封面只接受知乎和知图 HTTPS 域名；加载失败时使用稳定占位。")
    bullet(doc, "完整正文不得写入浏览器本地存储、代码仓库或演示素材。")

    heading(doc, "部署信息", 2)
    table(doc, ["项目", "信息", "交付检查"], [
        ["公网前端", "https://zhihucat-six.vercel.app/campaign", "发布最终构建后从无痕窗口访问"],
        ["前端代理", "/argus-api/:path*", "确认生产 rewrite 指向本轮后端"],
        ["后端健康检查", "部署后填写或通过代理访问 /health", "版本、时间和状态必须与最终构建一致"],
        ["演示备份", "离线案卷和录屏", "断开上游网络仍能完成主流程"],
    ], [1.35, 3.3, 2.4], font_size=8.65, first_col_bold=True)

    doc.add_page_break()
    heading(doc, "九 技术实现与交付边界", 1)
    table(doc, ["模块", "当前代码职责", "交付前必须确认"], [
        ["Next.js 前端", "入口、刘看山、调查、对决、End 和响应式样式", "最终合并分支无冲突标记并通过构建"],
        ["Node API", "知乎内容代理、案卷转换、对决与账户接口", "路由、CORS、限流和生产启动命令可用"],
        ["知乎内容客户端", "字段校验、八秒超时和错误归一化", "列表与详情各完成在线和失败测试"],
        ["案卷编排", "六条线索、四张牌规则、假说和求证问题", "内容负责人核对事实与转述标记"],
        ["本地进度", "记录完成分数和入口最佳成绩", "隐私模式失败不阻断当前局"],
        ["素材", "刘看山 GIF、静态图、证据图和音频", "路径、加载失败占位和授权说明完整"],
    ], [1.25, 3.25, 2.55], font_size=8.55, first_col_bold=True)

    heading(doc, "完成状态的定义", 2)
    status(doc, "可声明已完成", "功能在最终合并分支通过构建、自动化测试和一次在线加一次离线人工走查。", ui.GREEN)
    status(doc, "只可声明代码存在", "页面或路由已经编写，但当前环境尚未完成部署或端到端验证。", ui.ORANGE)
    status(doc, "不得声明", "多故事可玩、自由对话 Agent、完整原文阅读和原作结局还原。", ui.RED)

    heading(doc, "建议运行的验证命令", 2)
    table(doc, ["位置", "命令", "通过标准"], [
        ["仓库根目录", "npm test", "所有工作区测试通过"],
        ["前端", "npm run typecheck 和 npm run build", "无 TypeScript 或构建错误"],
        ["后端", "npm run typecheck 和 npm run build", "无冲突标记、类型或产物错误"],
        ["线上", "访问 /campaign 和代理 /health", "页面加载且后端版本为最终部署"],
    ], [1.2, 3.6, 2.25], font_size=8.65, first_col_bold=True)

    doc.add_page_break()
    heading(doc, "十 四人四小时执行计划", 1)
    paragraph(doc, "四人同时工作，但只有一名集成负责人可以合并主分支。每个人可以用 Astra 辅助搜索、编码、文案核对或测试设计；任何模型产出都必须由对应负责人核验后再进入最终构建。")

    heading(doc, "角色分工", 2)
    table(doc, ["角色", "主要责任", "Astra 使用", "四小时交付物"], [
        ["A 产品与内容负责人", "冻结范围、核对蓝血事实与边界、维护刘看山文案、组织演示", "检查文案一致性、生成备选表达，不补写事实", "产品口径、故事脚本、讲稿和提交说明"],
        ["B 前端与 UI 负责人", "完成页面状态、响应式、无障碍和刘看山交互", "定位组件问题、补测试和样式建议", "可操作前端、移动端修复和视觉走查"],
        ["C 后端与 API 负责人", "知乎接口、降级、路由、部署和健康检查", "分析日志、补接口测试和部署检查", "在线与离线接口、后端部署和版本证据"],
        ["D 集成与 QA 负责人", "控制合并、跑测试、端到端走查、录屏和提交", "生成测试矩阵、复核差异，不伪造测试通过", "最终分支、测试记录、演示备份和提交包"],
    ], [1.35, 2.45, 1.65, 1.6], font_size=8.1, first_col_bold=True)

    heading(doc, "相对时间表", 2)
    table(doc, ["时间", "全队门槛", "A 产品内容", "B 前端 UI", "C 后端 API", "D 集成 QA"], [
        ["T0 至 0:20", "冻结 P0 与分支", "锁定故事和文案边界", "确认页面缺口", "确认接口和部署目标", "建立合并顺序与测试清单"],
        ["0:20 至 1:20", "各自提交最小改动", "完成向导和演示文案", "修主流程和移动端阻塞", "修接口、降级和测试", "准备测试数据与设备"],
        ["1:20 至 2:10", "第一次集成构建", "逐页核对来源与措辞", "处理集成 UI 缺陷", "联调在线与离线路径", "合并并跑首轮 build test"],
        ["2:10 至 3:00", "两条路径走通", "按实际页面调整讲稿", "修 P0 交互和遮挡", "部署后端并查 health", "桌面与移动人工走查"],
        ["3:00 至 3:35", "代码冻结", "准备答辩问答", "只修阻断性问题", "只修阻断性问题", "部署前端、无痕验证、录屏"],
        ["3:35 至 4:00", "提交并保留备份", "主讲三分钟演示", "现场操作备份", "监控接口与降级", "核对提交项和保存证据"],
    ], [0.9, 1.3, 1.25, 1.2, 1.2, 1.25], font_size=7.55, first_col_bold=True)

    heading(doc, "四个硬门槛", 2)
    bullet(doc, "二十分钟门槛：只保留《蓝血》单故事，禁止临时加入第二故事或自由聊天。")
    bullet(doc, "两小时十分钟门槛：最终合并分支必须无冲突标记并至少完成一次构建；否则停止新增功能。")
    bullet(doc, "三小时门槛：在线和离线两条路径均能从入口走到 End；未通过时优先保证离线演示。")
    bullet(doc, "三小时三十五分钟门槛：冻结代码，只允许修复阻断演示、数据错误或来源误标。")

    heading(doc, "并行协作规则", 2)
    bullet(doc, "每项任务只设一名负责人；其他人提供审查，不同时改同一文件。")
    bullet(doc, "提交信息说明影响页面、测试方法和回退方式；集成负责人按依赖顺序合并。")
    bullet(doc, "Astra 可以提出补丁，但不能代替人工确认接口响应、版权边界和测试结果。")
    bullet(doc, "发现 P1 或 P2 问题时记录到清单，不挤占 P0 修复时间。")

    doc.add_page_break()
    heading(doc, "十一 验收指标与测试", 1)
    heading(doc, "P0 产品指标", 2)
    table(doc, ["指标", "通过标准", "验证方式"], [
        ["首屏理解", "一分钟内能说出故事冲突、玩家任务和内容边界", "让一名非项目成员复述"],
        ["完整路径", "入口到 End 可在约三分钟内完成", "在线和离线各走一遍"],
        ["玩法门槛", "六条线索、三张锚点、一张方向牌和四张出示规则全部生效", "按正常与错误路径操作"],
        ["来源准确", "在线命中、零命中和离线三种模式的文案与数据一致", "检查 mode 与 verifiedQuoteCount"],
        ["失败可恢复", "失败后返回调查且搜证进度不丢失", "主动制造失败再重组"],
        ["完成可判定", "未选假说或问题时不能完成，完成后分数保存", "分别测试禁用与成功路径"],
    ], [1.25, 3.5, 2.3], font_size=8.55, first_col_bold=True)

    heading(doc, "工程与界面验收", 2)
    for item in [
        "最终分支中不存在冲突标记、调试密钥或未解释的本地地址。",
        "前端和后端 typecheck、build 与相关测试均通过，并保存命令结果。",
        "1440、1024、390 和 360 宽度无横向滚动、遮挡或裁切。",
        "刘看山面板不遮挡主要按钮，移动端内部滚动和关闭可用。",
        "键盘可以完成入口、搜证、组牌、对决和 End 的主要操作。",
        "来源状态、选择、错误和完成均同时使用文字与非颜色提示。",
        "减少动态效果模式只显示静态素材，音频有可见静音入口。",
        "上游接口超时后进入透明降级，主流程仍可完成。",
    ]:
        bullet(doc, item, checkbox=True)

    heading(doc, "演示验收", 2)
    table(doc, ["场景", "必须展示", "备份"], [
        ["正常演示", "刘看山导览、六线索、四卡对决、End 求证", "预先打开页面并保留一局可重玩状态"],
        ["接口在线", "来源徽标、Work ID 和至少一个实际命中状态", "若零命中则如实展示未命中"],
        ["接口异常", "策划离线模式和完整主流程", "断网上游或使用已验证的离线案卷"],
        ["部署故障", "本地构建或录屏", "保留三分钟无剪辑走查视频"],
    ], [1.25, 3.85, 1.95], font_size=8.65, first_col_bold=True)

    heading(doc, "完成定义", 2)
    paragraph(doc, "只有 P0 指标全部通过、在线和离线各完成一次人工走查、最终公网地址加载的是同一份冻结构建，并且提交材料中的能力描述与实际页面一致，项目才能标记为可交付。任何未验证能力保留为待完成。")

    heading(doc, "十二 风险登记表", 1)
    table(doc, ["风险", "概率", "影响", "触发信号", "处理与负责人"], [
        ["知乎接口不可用或字段变化", "高", "高", "超时、非 2xx、字段校验失败", "进入 curated fallback；C 负责日志与恢复，A 核对页面措辞"],
        ["把转述误标成原文", "中", "高", "verifiedQuoteCount 与页面不一致", "只由后端命中生成标记；A 负责内容审查"],
        ["合并后无法构建", "中", "高", "冲突标记、类型错误、测试失败", "D 控制合并，2:10 后停止新增功能"],
        ["线上仍是旧构建", "中", "高", "health 版本或页面与本地不一致", "C 验后端，D 用无痕窗口验证前端并记录时间"],
        ["移动端遮挡或双滚动", "中", "中", "390 或 360 宽度无法完成操作", "B 按设备矩阵修复，D 实机复验"],
        ["规则难懂导致卡关", "中", "中", "玩家不知缺少哪条线索或牌", "B 显示缺失项，A 调整刘看山提示"],
        ["自由生成内容产生幻觉", "低", "高", "回答没有材料 ID 或越过边界", "本轮不开放自由问答；未来失败即回退预设"],
        ["远程图片或音频失败", "低", "中", "封面空白、音频叠播", "B 使用占位和静态素材，D 检查退出清理"],
        ["演示时间超出三分钟", "中", "中", "搜证或对决讲解过长", "A 使用固定讲稿，D 录制计时备份"],
    ], [1.45, 0.55, 0.55, 2.05, 2.55], font_size=7.85, first_col_bold=True)

    heading(doc, "风险处理顺序", 2)
    bullet(doc, "先处理来源误标、无法构建和主流程中断，这三类问题会直接使产品不可交付。")
    bullet(doc, "再处理移动端和规则可读性；没有时间时保留可用布局，不追加装饰动效。")
    bullet(doc, "接口异常优先启用透明降级，不在演示现场临时修改上游数据。")

    doc.add_page_break()
    heading(doc, "十三 三分钟演示讲稿", 1)
    table(doc, ["时间", "操作", "讲解要点"], [
        ["0:00 至 0:20", "打开公网入口", "这是《蓝血疑云》，我们把知乎故事公开片段变成一局可质证的互动案卷。页面先说明来源状态和改编边界。"],
        ["0:20 至 0:45", "点击先听刘看山讲故事", "刘看山不是故事角色，也不是自由聊天机器人。他先讲清冲突、玩法和目前不能确定的部分。"],
        ["0:45 至 1:25", "进入调查并快速收集六条线索", "三处场景把身体异常、公共记录和被观察的迹象分开。右侧材料会显示原文短引、未命中或策划转述。"],
        ["1:25 至 1:55", "选择三张锚点和一张方向牌", "关键事实固定入选，第四张决定我们更重视世界规则变化还是主动观察。背景教材不能上庭。"],
        ["1:55 至 2:25", "进入对决并出示四张证据", "玩家不是点最高分牌，而是让四项材料形成证据链，削弱疲劳、误记和巧合的解释。"],
        ["2:25 至 2:50", "进入 End 选择假说和求证问题", "系统不宣布真相，只比较支持与缺口。推荐问题兼顾信息价值和暴露风险，把下一步带回知乎。"],
        ["2:50 至 3:00", "指出来源状态与降级", "即使接口临时不可用，玩法仍完整，但页面会明确说这是策划离线转述，不冒充已核验原文。"],
    ], [0.85, 2.2, 4.0], font_size=8.35, first_col_bold=True)

    heading(doc, "答辩常见问题", 2)
    table(doc, ["问题", "建议回答"], [
        ["为什么只做一个故事", "四小时内优先证明一条稳定闭环。案卷结构可以复用，但其他故事尚未完成编排，不能宣称可玩。"],
        ["知乎 API 的价值在哪里", "API 决定故事目录、作者、简介、来源模式和短引核验；这些状态直接进入证据阅读器和 End。"],
        ["为什么不用自由对话 Agent", "当前演示需要可重复和可核验。刘看山先用确定性答案完成 P0，自由问答留到来源约束完善后。"],
        ["是不是改写了原作结局", "没有。所有假说和结果都标为阶段推演，完整正文不落盘，玩家最后选择的是下一步验证问题。"],
        ["接口挂了怎么办", "后端返回内置案卷，页面切到策划离线模式并清除原文核验标记，主流程继续。"],
    ], [2.05, 5.0], font_size=8.7, first_col_bold=True)

    heading(doc, "演示前十分钟", 2)
    for item in [
        "用无痕窗口打开 https://zhihucat-six.vercel.app/campaign，并确认入口不是旧构建。",
        "检查代理 health、知乎故事列表和《蓝血》案卷详情；记录当前来源模式。",
        "完整走一遍演示路径并计时，确保不使用尚未实现的功能。",
        "关闭不相关标签页和通知，准备离线案卷、录屏和本地运行三个备份。",
        "确认演示账号、网络、电源和声音；音频异常时直接静音，不影响主讲。",
    ]:
        bullet(doc, item, checkbox=True)

    heading(doc, "十四 提交清单", 1, page_break=True)
    heading(doc, "必须提交", 2)
    for item in [
        "公网演示地址 https://zhihucat-six.vercel.app/campaign 可在无痕窗口访问。",
        "代码仓库使用最终冻结提交，README 写明启动方式、环境变量、测试命令和内容边界。",
        "三分钟演示视频或现场演示备份包含入口、刘看山、搜证、组牌、对决、End 和降级说明。",
        "产品说明计划书和 UI 实现说明均放入 docs 目录。",
        "团队成员、角色分工、项目简介、使用技术和知乎 API 用法填写完整。",
        "原作作者、Work ID、查看原作入口和改编声明清楚可见。",
        "前端与后端测试结果、构建结果和人工走查记录保存在提交说明中。",
        "环境变量与密钥未提交到仓库，演示账号不包含个人敏感信息。",
    ]:
        bullet(doc, item, checkbox=True)

    heading(doc, "提交信息", 2)
    table(doc, ["项目", "内容"], [
        ["产品名称", "蓝血疑云"],
        ["公网演示", "https://zhihucat-six.vercel.app/campaign"],
        ["代码仓库", "https://github.com/988ms/zhihucat"],
        ["知乎故事", "《蓝血》 桃花先生 Work ID 2025684191967294692"],
        ["后端健康地址", "部署后填写或通过前端代理 /health 访问"],
        ["演示视频", "录制后填写"],
        ["版本提交", "代码冻结后填写 commit SHA"],
    ], [1.45, 5.6], font_size=8.8, first_col_bold=True)

    heading(doc, "不得在提交材料中声称", 2)
    bullet(doc, "已经支持多个知乎故事完整游玩。")
    bullet(doc, "已经读取、保存或展示知乎故事完整正文。")
    bullet(doc, "刘看山已经接入自由对话大模型。")
    bullet(doc, "阶段假说等同于原作结局或作者答案。")
    bullet(doc, "尚未运行的测试、尚未部署的路由或尚未真机验收的界面已经通过。")

    heading(doc, "最终签字检查", 2)
    table(doc, ["负责人", "确认内容", "结果"], [
        ["A 产品与内容", "故事事实、来源标识、讲稿和提交描述一致", "待确认"],
        ["B 前端与 UI", "主流程、移动端、键盘和素材状态可用", "待确认"],
        ["C 后端与 API", "在线、零命中、离线、部署和 health 可用", "待确认"],
        ["D 集成与 QA", "最终分支、测试记录、演示地址和备份一致", "待确认"],
    ], [1.55, 4.5, 1.0], font_size=8.75, first_col_bold=True)
    paragraph(doc, "最终判定  四名负责人只对已经验证的结果签字。若公网环境与本地结果不一致，以公网无痕窗口的实际表现为准，并在提交说明中如实记录降级状态。", bold_lead="最终判定")

    props = doc.core_properties
    props.title = "蓝血互动叙事游戏产品说明计划书"
    props.subject = "知乎黑客松蓝血互动叙事游戏的产品范围 实施计划 验收与演示说明"
    props.author = "蓝血互动案卷项目组"
    props.keywords = "蓝血, 刘看山, 知乎黑客松, 产品计划, 互动叙事, 知乎 API"
    props.comments = "区分已存在代码 待联调和尚未实现能力"

    for item in doc.tables:
        item.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in item.rows:
            ui.set_no_cell_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    ui.set_paragraph_keep(p)

    doc.save(str(OUTPUT))


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
